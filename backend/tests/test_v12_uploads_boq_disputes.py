"""Uploads (logo, product image, documents), BOQ import, bid comparison export, disputes."""
import io
import struct
import zlib

import openpyxl
from conftest import auth

API = "/api/v1"


def _png() -> bytes:
    def chunk(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b"\x00\xff\x00\x00\xff"
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def test_uploads_and_documents(client):
    sup = auth(client, "supplier3@demo.sa")
    admin = auth(client, "admin@mysupplier.sa", "Admin@2026")
    r = client.post(f"{API}/suppliers/me/logo", headers=sup, files={"file": ("logo.png", io.BytesIO(_png()), "image/png")})
    assert r.status_code == 200 and "/uploads/logos/" in r.json()["logo_url"]
    assert client.get(r.json()["logo_url"].replace("http://localhost:8000", "")).status_code == 200
    bad = client.post(f"{API}/suppliers/me/logo", headers=sup, files={"file": ("x.txt", io.BytesIO(b"hello"), "text/plain")})
    assert bad.status_code == 400
    r = client.post(f"{API}/suppliers/me/documents", headers=sup, params={"kind": "cr"}, files={"file": ("cr.pdf", io.BytesIO(b"%PDF-1.4 test"), "application/pdf")})
    assert r.status_code == 201 and r.json()["status"] == "pending"
    doc_id = r.json()["id"]
    assert client.get(f"{API}/suppliers/me/documents", headers=sup).json()[0]["id"] == doc_id
    pending = client.get(f"{API}/admin/documents", headers=admin).json()
    assert any(d["id"] == doc_id and d["supplier_name"] for d in pending)
    r = client.post(f"{API}/admin/documents/{doc_id}/review", headers=admin, json={"status": "approved", "note": "ok"})
    assert r.json()["status"] == "approved" and r.json()["reviewed_at"]
    # product image by creator
    cats = client.get(f"{API}/catalog/categories").json()
    p = client.post(f"{API}/catalog/products", headers=sup, json={"category_id": cats[0]["id"], "name_ar": "منتج بصورة", "name_en": "Pictured", "unit": "piece"}).json()
    r = client.post(f"{API}/catalog/products/{p['id']}/image", headers=sup, files={"file": ("p.png", io.BytesIO(_png()), "image/png")})
    assert r.status_code == 200 and "/uploads/products/" in r.json()["image_url"]
    assert client.post(f"{API}/catalog/products/{p['id']}/image", headers=auth(client, "supplier1@demo.sa"), files={"file": ("p.png", io.BytesIO(_png()), "image/png")}).status_code == 403


def test_boq_import_and_comparison_export(client):
    buyer = auth(client, "buyer@demo.sa")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["مشروع فيلا — جدول الكميات"])
    ws.append(["م", "البند", "الوحدة", "الكمية", "السعر"])
    ws.append([1, "حديد تسليح قطر 12 مم", "طن", 25, 2500])
    ws.append([2, "أسمنت بورتلاندي عادي 50 كجم", "كيس", 800, None])
    ws.append([3, "بند غير معروف تماماً xyz", "م2", 40, None])
    ws.append([4, "", "", None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"{API}/rfq/import-boq", headers=buyer, files={"file": ("boq.xlsx", io.BytesIO(buf.getvalue()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    items = r.json()
    assert len(items) == 3
    assert items[0]["product_id"] and "12" in items[0]["match_name_ar"] and items[0]["quantity"] == 25 and items[0]["target_price"] == 2500
    assert items[1]["product_id"] and items[2]["product_id"] is None
    # CSV variant
    csv_data = "description,qty,unit\nRebar 16mm,10,ton\n"
    r = client.post(f"{API}/rfq/import-boq", headers=buyer, files={"file": ("boq.csv", io.BytesIO(csv_data.encode()), "text/csv")})
    assert r.status_code == 200 and r.json()[0]["product_id"]
    # create RFQ from items, get two bids, export comparison
    rfq = client.post(f"{API}/rfq", headers=buyer, json={"title": "من جدول الكميات", "city": "الرياض",
                      "items": [{"product_id": i["product_id"], "description": i["description"], "quantity": i["quantity"], "unit": i["unit"]} for i in items]}).json()
    for email, base in (("supplier1@demo.sa", 100), ("supplier2@demo.sa", 90)):
        h = auth(client, email)
        assert client.post(f"{API}/rfq/{rfq['id']}/bids", headers=h, json={"items": [{"rfq_item_id": it["id"], "unit_price": base + k} for k, it in enumerate(rfq["items"])]}).status_code == 201
    r = client.get(f"{API}/rfq/{rfq['id']}/export.xlsx", headers=buyer)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb2 = openpyxl.load_workbook(io.BytesIO(r.content))
    ws2 = wb2.active
    assert ws2.max_row >= 3 + 7 and ws2.cell(row=2, column=5).value  # supplier column present
    assert client.get(f"{API}/rfq/{rfq['id']}/export.xlsx", headers=auth(client, "supplier1@demo.sa")).status_code == 404


def test_disputes(client):
    buyer = auth(client, "buyer@demo.sa")
    admin = auth(client, "admin@mysupplier.sa", "Admin@2026")
    p = client.get(f"{API}/catalog/products", params={"q": "طوب أحمر"}).json()["items"][0]
    offer = next(o for o in client.get(f"{API}/catalog/products/{p['id']}").json()["offers"] if not o["supplier"]["is_external"])
    order = client.post(f"{API}/orders/direct", headers=buyer, json={"offer_id": offer["id"], "quantity": max(offer["min_qty"], 1000)}).json()
    pay = client.post(f"{API}/payments/checkout", headers=buyer, json={"order_id": order["id"]}).json()
    client.post(f"{API}/payments/mock/confirm/{pay['id']}", data={"result": "paid"}, follow_redirects=False)
    r = client.post(f"{API}/orders/{order['id']}/dispute", headers=buyer, json={"reason": "المواد وصلت تالفة ولم يتم استبدالها"})
    assert r.status_code == 201 and r.json()["status"] == "open" and r.json()["role"] == "buyer"
    assert client.post(f"{API}/orders/{order['id']}/dispute", headers=buyer, json={"reason": "duplicate dispute"}).status_code == 409
    open_list = client.get(f"{API}/admin/disputes", headers=admin).json()
    d = next(x for x in open_list if x["order_id"] == order["id"])
    assert d["supplier_name"] and d["order_total"] == order["total"]
    r = client.post(f"{API}/admin/disputes/{d['id']}/resolve", headers=admin, json={"status": "resolved", "resolution": "رد كامل المبلغ", "refund": True})
    assert r.json()["status"] == "resolved" and r.json()["refunded"] is True
    o = client.get(f"{API}/orders/{order['id']}", headers=buyer).json()
    assert o["status"] == "cancelled" and o["payment_status"] == "refunded"
    assert client.get(f"{API}/orders/{order['id']}/disputes", headers=buyer).json()[0]["status"] == "resolved"
