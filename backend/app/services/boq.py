"""Bill-of-quantities import (Excel/CSV) → RFQ items, with fuzzy product matching; bid comparison export."""
import csv
import io
import re
from difflib import SequenceMatcher

from sqlalchemy.orm import Session

from ..models import RFQ, Product

_DESC = ("description", "item", "material", "الوصف", "البند", "بيان", "المادة", "الصنف", "name")
_QTY = ("qty", "quantity", "الكمية", "كمية")
_UNIT = ("unit", "uom", "الوحدة", "وحدة")
_PRICE = ("rate", "unit price", "price", "السعر", "سعر الوحدة")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _find_col(headers: list[str], names: tuple[str, ...]) -> int | None:
    for i, h in enumerate(headers):
        for n in names:
            if n in h:
                return i
    return None


def _rows_from_bytes(data: bytes, filename: str) -> list[list]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def _to_float(v) -> float | None:
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def parse_boq(data: bytes, filename: str) -> list[dict]:
    rows = [r for r in _rows_from_bytes(data, filename) if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    # header row = first row that contains a description-like and a quantity-like header
    header_idx = 0
    for i, r in enumerate(rows[:15]):
        hs = [_norm(c) for c in r]
        if _find_col(hs, _DESC) is not None and _find_col(hs, _QTY) is not None:
            header_idx = i
            break
    headers = [_norm(c) for c in rows[header_idx]]
    ci = _find_col(headers, _DESC) if _find_col(headers, _DESC) is not None else 0
    cq = _find_col(headers, _QTY)
    cu = _find_col(headers, _UNIT)
    cp = _find_col(headers, _PRICE)
    out = []
    for r in rows[header_idx + 1:]:
        desc = str(r[ci]).strip() if ci < len(r) and r[ci] is not None else ""
        if not desc:
            continue
        qty = _to_float(r[cq]) if cq is not None and cq < len(r) else None
        if qty is None:  # look for the first numeric cell after the description
            qty = next((_to_float(c) for c in r[ci + 1:] if _to_float(c) is not None), None)
        if not qty or qty <= 0:
            continue
        out.append({"description": desc, "quantity": qty,
                    "unit": str(r[cu]).strip() if cu is not None and cu < len(r) and r[cu] else "",
                    "target_price": _to_float(r[cp]) if cp is not None and cp < len(r) else None})
    return out[:300]


def match_products(db: Session, items: list[dict]) -> list[dict]:
    products = db.query(Product).filter(Product.is_active.is_(True)).all()
    for it in items:
        d = _norm(it["description"])
        best, score = None, 0.0
        for p in products:
            for cand in (p.name_ar, p.name_en, f"{p.name_ar} {p.brand}", p.sku):
                s = SequenceMatcher(None, d, _norm(cand)).ratio()
                tokens = set(d.split()) & set(_norm(cand).split())
                s = max(s, len(tokens) / max(1, len(set(_norm(cand).split()))) * 0.9)
                if s > score:
                    best, score = p, s
        if best and score >= 0.45:
            it.update({"product_id": best.id, "match_name_ar": best.name_ar, "match_name_en": best.name_en, "confidence": round(score, 2)})
            if not it["unit"]:
                it["unit"] = best.unit
        else:
            it.update({"product_id": None, "match_name_ar": "", "match_name_en": "", "confidence": round(score, 2)})
    return items


def bid_comparison_xlsx(rfq: RFQ) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.sheet_view.rightToLeft = True
    bids = sorted([b for b in rfq.bids if b.status != "withdrawn"], key=lambda b: b.total)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="175934")
    best_fill = PatternFill("solid", fgColor="EDF5EF")
    ws.append([f"{rfq.title} — RFQ #{rfq.id}", "", rfq.city])
    ws["A1"].font = Font(bold=True, size=13)
    headers = ["البند / Item", "الكمية", "الوحدة", "مرجع السوق (أدنى)"] + [b.supplier.name for b in bids]
    ws.append(headers)
    for c in ws[2]:
        c.font, c.fill, c.alignment = head, fill, Alignment(horizontal="center", wrap_text=True)
    for it in rfq.items:
        prices = []
        for b in bids:
            bi = next((x for x in b.items if x.rfq_item_id == it.id), None)
            prices.append(bi.unit_price if bi else None)
        ws.append([it.description, it.quantity, it.unit, None] + prices)
        row = ws.max_row
        valid = [p for p in prices if p is not None]
        if valid:
            for j, p in enumerate(prices):
                if p == min(valid):
                    ws.cell(row=row, column=5 + j).fill = best_fill
    ws.append([])
    ws.append(["المجموع قبل الضريبة", "", "", ""] + [b.subtotal for b in bids])
    ws.append(["الضريبة 15%", "", "", ""] + [b.vat for b in bids])
    ws.append(["الإجمالي", "", "", ""] + [b.total for b in bids])
    ws.append(["مدة التوريد (يوم)", "", "", ""] + [b.delivery_days for b in bids])
    ws.append(["الترتيب", "", "", ""] + list(range(1, len(bids) + 1)))
    ws.append(["الحالة", "", "", ""] + [b.status for b in bids])
    for r in ws.iter_rows(min_row=ws.max_row - 5, max_row=ws.max_row):
        r[0].font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    ws.column_dimensions["A"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
